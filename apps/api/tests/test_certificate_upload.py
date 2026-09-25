import io
import unittest

from docx import Document
from fastapi import UploadFile
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.db import Base
from app.models import User
from app.routers.career import _ingest_upload, list_documents


class CertificateUploadTest(unittest.IsolatedAsyncioTestCase):
    async def test_docx_and_pdf_certificates_are_indexed_and_listed(self):
        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        try:
            async with engine.begin() as connection:
                await connection.run_sync(Base.metadata.create_all)

            session_factory = async_sessionmaker(engine, expire_on_commit=False)
            async with session_factory() as db:
                user = User(id="certificate-test-user", email="certificate-test@example.com", name="Test User")
                db.add(user)
                await db.commit()

                document = Document()
                document.add_paragraph("Certificate of Completion")
                document.add_paragraph("Awarded for completing the data analysis course.")
                file_bytes = io.BytesIO()
                document.save(file_bytes)
                file_bytes.seek(0)

                upload = UploadFile(file=file_bytes, filename="certificate.docx")
                result = await _ingest_upload(file=upload, document_type="certificate", user=user, db=db)
                self.assertTrue(result["ok"])
                self.assertEqual(result["document"]["source_type"], "certificate")
                self.assertGreater(result["document"]["chunk_count"], 0)

                saved = await list_documents(user=user, db=db)
                self.assertEqual(len(saved), 1)
                self.assertEqual(saved[0]["id"], result["document"]["id"])
                self.assertEqual(saved[0]["status"], "ready")

                pdf_bytes = io.BytesIO()
                pdf = Canvas(pdf_bytes)
                pdf.drawString(72, 720, "Certificate of Completion")
                pdf.save()
                pdf_bytes.seek(0)
                pdf_upload = UploadFile(file=pdf_bytes, filename="certificate.pdf")
                pdf_result = await _ingest_upload(file=pdf_upload, document_type="certificate", user=user, db=db)
                self.assertTrue(pdf_result["ok"])
                self.assertGreater(pdf_result["document"]["chunk_count"], 0)
                self.assertEqual(len(await list_documents(user=user, db=db)), 2)
        finally:
            await engine.dispose()


if __name__ == "__main__":
    unittest.main()
